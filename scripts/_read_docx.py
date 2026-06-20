from docx import Document

print("=== QA_Document.docx Q32-Q34 context ===")
doc = Document('QA_Document.docx')
for i in range(83, 100):
    if i < len(doc.paragraphs):
        t = doc.paragraphs[i].text.strip()
        safe = t.encode('ascii', 'replace').decode()
        print(f'[{i}] {safe[:220]}')

print()
print("=== Streamlit_Output_Explained.docx paras 63-80 ===")
doc2 = Document('Streamlit_Output_Explained.docx')
for i in range(63, 80):
    if i < len(doc2.paragraphs):
        t = doc2.paragraphs[i].text.strip()
        safe = t.encode('ascii', 'replace').decode()
        print(f'[{i}] {safe[:220]}')
