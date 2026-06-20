"""Insert literature comparison section into Battery_Project_Report.docx before Limitations."""
from docx import Document
from docx.oxml.ns import qn

doc = Document("Battery_Project_Report.docx")

# Find the Limitations heading paragraph
lim_heading_el = None
for p in doc.paragraphs:
    if p.text.strip() == "14. Limitations and Honest Assessment":
        lim_heading_el = p._p
        break

if lim_heading_el is None:
    print("ERROR: Limitations heading not found")
    exit(1)

def make_para(doc, text, bold_heading=False):
    p = doc.add_paragraph(text)
    if bold_heading:
        for run in p.runs:
            run.bold = True
    return p

# Build paragraphs to insert
paras_to_insert = []

# Section heading
h = doc.add_paragraph("14. Comparison with Prior Work")
try:
    h.style = doc.styles["Heading 1"]
except:
    for run in h.runs:
        run.bold = True
paras_to_insert.append(h)

# Caveat paragraph
paras_to_insert.append(doc.add_paragraph(
    "IMPORTANT: Direct RMSE comparison with prior work is misleading because the tasks differ "
    "fundamentally. Within-battery methods (Zhao 2022, Catelani 2021) observe early cycles of "
    "the same target battery and extrapolate forward -- the target cell is partially visible. "
    "Cross-battery methods generalise to entirely unseen cells. The latter is harder and "
    "produces higher RMSE by design. The comparison is included for completeness only."
))

# Table header as text rows (python-docx table insertion mid-doc is complex; use structured text)
paras_to_insert.append(doc.add_paragraph(
    "Paper | Dataset | Task type | EOL definition | Reported error | Uncertainty | Multi-seed"
))
paras_to_insert.append(doc.add_paragraph(
    "Lin et al. 2023 (GRU+HMM) | Oxford + NASA RW | SOH curve fitting | SOH % threshold | "
    "RMSE 0.2-1.8% SOH | HMM residual bounds | No"
))
paras_to_insert.append(doc.add_paragraph(
    "Zhao et al. 2022 (BLS-LSTM) | NASA B0005/B0006 + CALCE CX2 | Within-battery RUL | "
    "70% of nominal | AE approx 1 cycle | None | No"
))
paras_to_insert.append(doc.add_paragraph(
    "Catelani et al. 2021 (ESN) | NASA B0005-B0018 | Within-battery RUL | 70% of nominal | "
    "EE 0-6 cycles | t-distribution bounds | No"
))
paras_to_insert.append(doc.add_paragraph(
    "THIS WORK (XGBoost + TCN ensemble) | NASA PCoE 34 cells, 3 thermal groups | "
    "Cross-battery RUL | 80% of per-cell measured init cap (0.81-1.53 Ah range) | "
    "RMSE 44.74 +/- 8.38 cycles (5 seeds) | Split-conformal per temperature group | Yes (5 seeds)"
))

paras_to_insert.append(doc.add_paragraph(
    "Key differentiators: (1) Cross-battery task -- target cell entirely unseen at test time; "
    "(2) Split-conformal uncertainty with coverage guarantee; (3) Per-group calibration "
    "(LOBO for cold, split for room/hot); (4) Coverage failure mechanistically linked to "
    "PSI drift; (5) 5-seed reproducibility with reported std. "
    "Prior works use single-seed, within-battery setups without conformal uncertainty."
))

# Update section numbering -- old 14 (Limitations) becomes 15, old 15 (Equations) becomes 16
for p in doc.paragraphs:
    if p.text.strip() == "14. Limitations and Honest Assessment":
        for run in p.runs:
            run.text = run.text.replace("14.", "15.")
        if not p.runs:
            p.add_run("15. Limitations and Honest Assessment")
        break

for p in doc.paragraphs:
    if p.text.strip() == "15. Key Equations Reference":
        for run in p.runs:
            run.text = run.text.replace("15.", "16.")
        if not p.runs:
            p.add_run("16. Key Equations Reference")
        break

# XML-insert all new paragraphs before old Limitations heading
for np_ in reversed(paras_to_insert):
    lim_heading_el.addprevious(np_._p)

doc.save("Battery_Project_Report.docx")
print(f"Inserted {len(paras_to_insert)} paragraphs. Sections renumbered 14->15, 15->16.")
