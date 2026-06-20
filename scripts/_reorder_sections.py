"""Reorder reversed sections: indices 207-223 resequenced correctly."""
from docx import Document

doc = Document("Battery_Project_Report.docx")
body = doc.element.body
P = doc.paragraphs

# Correct order of paragraph indices (207-223 scrambled by prior insertions)
# Current wrong order: 207=lim7, 208=lim6, 209=lim5, 210=lim4, 211=lim3, 212=lim2, 213=lim1
#                      214=key_diff, 215=thiswork, 216=catelani, 217=zhao, 218=lin, 219=header
#                      220=caveat, 221=lit_heading, 222=lim_heading, 223=key_eq
# Correct order: lit_heading, caveat, header, lin, zhao, catelani, thiswork, key_diff,
#                lim_heading, lim1, lim2, lim3, lim4, lim5, lim6, lim7, key_eq

correct_indices = [221, 220, 219, 218, 217, 216, 215, 214,
                   222, 213, 212, 211, 210, 209, 208, 207,
                   223]

# Collect the paragraph XML elements in current order (207..223)
section_els = [P[i]._p for i in range(207, 224)]

# Remove them all from body
for el in section_els:
    body.remove(el)

# Now append them in correct order
for idx in correct_indices:
    # idx refers to original paragraph index before removal
    body.append(section_els[idx - 207])

doc.save("Battery_Project_Report.docx")
print("Reordered. Verifying tail...")

doc2 = Document("Battery_Project_Report.docx")
for i, p in enumerate(doc2.paragraphs):
    if i >= 204:
        safe = p.text.strip().encode("ascii","replace").decode()
        print(f"[{i}] {safe[:90]}")
