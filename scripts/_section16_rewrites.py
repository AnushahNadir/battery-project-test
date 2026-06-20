"""Section 16 rewrites: Battery_Project_Report.docx full paper section corrections."""
from docx import Document

def set_para(para, text):
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)

def find_para(paras, fragment):
    for i, p in enumerate(paras):
        if fragment in p.text:
            return i
    return None

doc = Document("Battery_Project_Report.docx")
P = doc.paragraphs
changes = []
not_found = []

def fix(fragment, new_text, label=""):
    i = find_para(P, fragment)
    if i is not None:
        set_para(P[i], new_text)
        changes.append(f"[{i}] {label or fragment[:60]}")
    else:
        not_found.append(f"NOT FOUND: {fragment[:70]}")

# SECTION 5.1 Statistical baseline RMSE
fix(
    "Result: RMSE = 466.82 cycles",
    "Result: RMSE = 45.64 cycles (seed 42); multi-seed mean 319.66 +/- 196.58 cycles across 5 random battery splits (highly unstable). Auto-excluded from ensemble: RMSE far exceeds 1.5x best model threshold.",
    "stat baseline RMSE 466.82",
)

# SECTION 5.2 XGBoost RMSE
fix(
    "Result: RMSE = 22.08 cycles, MAE = 13.18 cycles",
    "Result (seed 42): RMSE = 21.74 cycles, MAE = 10.96 cycles. Multi-seed mean (5 seeds): RMSE = 44.74 +/- 8.38 cycles. The large gap between best seed (21.74) and multi-seed mean (44.74) reflects battery-split dependence. The multi-seed figure is the honest headline number.",
    "XGBoost RMSE 22.08->21.74 multi-seed",
)

# SECTION 5.3 TCN RMSE
fix(
    "Result: RMSE = 25.63 cycles",
    "Result (seed 42): RMSE = 27.35 cycles. Multi-seed mean: RMSE = 41.34 +/- 4.78 cycles (5 seeds). The DL model is more stable across seeds than XGBoost (std 4.78 vs 8.38).",
    "TCN RMSE 25.63->27.35 multi-seed",
)

# SECTION 5.4 Ensemble - two models not three
fix(
    "the three models are combined using inverse-RMSE weighting",
    "surviving models are combined using inverse-RMSE weighting. Any model whose RMSE exceeds 1.5x the best RMSE is auto-excluded. The statistical baseline is always excluded (highly unstable across seeds), leaving a two-model ensemble of XGBoost and TCN/MLP:",
    "ensemble three->two",
)
fix(
    "Statistical baseline: 0.00  (excluded",
    "Statistical baseline: 0.00  (excluded -- RMSE >> 1.5x best model threshold)",
    "stat weight line",
)
fix("XGBoost:              0.537", "XGBoost:              0.557  (ml_weight)", "xgb 0.537->0.557")
fix("TCN:                  0.463", "TCN/MLP:              0.443  (dl_weight)", "tcn 0.463->0.443")

# SECTION 5.5 CV/test ratio
fix(
    "The CV RMSE (48.17) is 2.2",
    "5-fold GroupKFold CV gives mean RMSE = 17.64 +/- 13.73 cycles (worst fold: 43.95 cycles). Multi-seed test RMSE = 44.74 +/- 8.38 cycles. The CV mean (17.64) is optimistic relative to test (44.74) because CV folds share the same battery-split seed. This is why the system receives CONDITIONAL PASS: results depend substantially on which batteries land in the test set.",
    "CV 48.17->17.64",
)

# SECTION 6 90%->80%
fix(
    "I am 90% confident",
    "Predicting RUL without a confidence interval is not useful in practice -- you need to know not just this battery has 50 cycles left but also I am 80% confident the true value lies between 30 and 70 cycles. This section describes calibrated prediction intervals using split-conformal prediction with an 80% coverage target.",
    "uncertainty intro 90%->80%",
)

# SECTION 6.2 alpha formula
fix(
    "1 - 0.90 = 0.10",
    "alpha = 1 - 0.80 = 0.20 (targeting 80% marginal coverage).",
    "alpha 0.90->0.80",
)

# SECTION 6.2 coverage guarantee
fix(
    "at least 90% of the true RUL values will fall inside",
    "The marginal guarantee: at least 80% of true RUL values will fall inside [RUL_lower, RUL_upper] -- provided calibration and test data are exchangeable. Empirically: overall 91.85% (above target). Per group -- Room: 100.0% (q_hat=69.71 cycles, split conformal, over-wide); Hot: 100.0% (q_hat=25.28 cycles); Cold: 74.5% (q_hat=31.98 cycles, LOBO) -- BELOW the 80% target. Mean CI width: 60.62 cycles. Cold-group failure is the central finding of this section.",
    "coverage guarantee 90%->80%",
)
fix(
    "Empirically we achieved 96.39%",
    "Empirically, overall conformal coverage was 91.85% (above the 80% target). Per-group: Room 100%, Hot 100%, Cold 74.5% (below target). Note: the 96.3-96.4% figure sometimes cited is raw [5th,95th]-percentile prediction coverage -- this is NOT the conformal coverage metric and must not be conflated.",
    "empirical 96.39%->91.85%",
)

# SECTION 6.4 Figure 7 caption
fix(
    "Cold remains slightly below",
    "Figure 7 -- Conformal coverage by temperature group. Room and hot: 100% (over-wide intervals). Cold (LOBO calibration): 74.5% empirical vs 80% target -- 5.5 pp below target. Under-coverage is mechanistically explained by RED-level PSI drift on 5/10 features (PSI range 1.01-2.57), which violates the exchangeability assumption underlying conformal guarantees.",
    "fig7 slightly below->5.5pp",
)

# SECTION 7 EOL 1.6 Ah
fix(
    "the EOL threshold (1.6 Ah)",
    "the per-cell EOL threshold (80% of each cell measured initial capacity; range across 34 cells: 0.81-1.53 Ah, mean 1.23 Ah). This is NOT a fixed 1.6 Ah cutoff -- it varies per cell because measured initial capacities range from 1.02 to 1.92 Ah.",
    "EOL 1.6 Ah->per-cell",
)

# SECTION 7 survival exploratory
fix(
    "survival analysis answers a different",
    "survival analysis answers a complementary question: what is the probability the battery is still operating at cycle t? The pipeline uses a discrete-time logistic hazard model (logistic-regression approximation to Cox PH). Critical limitation: only 5 EOL events were observed in the test set (event rate 0.78%), making survival metrics unreliable. No C-index is reported. Survival output is exploratory only.",
    "survival 5 events no C-index",
)

# SECTION 7.4 risk distribution
fix(
    "Result on the test set (638 rows across 6 batteries):",
    "Result: risk tier distribution across test-set prediction records (uncertainty_metrics.json, 2790 total records):",
    "risk header 638->2790",
)
fix(
    "LOW:    689 rows (24.7%)",
    "LOW:    212 observations (7.6%) -- early-life cycles",
    "LOW 689->212",
)
fix(
    "MEDIUM: 2,020 rows (72.4%)",
    "MEDIUM: 537 observations (19.2%) -- mid-life cycles",
    "MEDIUM 2020->537",
)
fix(
    "HIGH:   81 rows (2.9%)",
    "HIGH:   2,041 observations (73.2%) -- late-life cycles approaching EOL. High-risk dominance reflects the preponderance of late-life observations in the dataset.",
    "HIGH 81->2041",
)

# SECTION 8 anomaly unsupervised
fix(
    "Result: 32 anomalies detected across the 6 test batteries.",
    "Result: 32 cycles flagged across the 6 test batteries using unsupervised residual-threshold detection (cycles where |RUL_true - RUL_predicted| exceeds the 95th-percentile training residual). No ground-truth anomaly labels exist; precision and recall cannot be computed. Zero flagged is also valid -- it means the battery degraded smoothly.",
    "anomaly unsupervised no P/R",
)

# SECTION 9 drift
fix(
    "Figure 9",
    "Figure 9 -- PSI drift report: train vs test. Five features show RED-level drift (PSI > 0.20): i_mean (2.573), temp_max (2.462), v_mean (2.077), duration_s (1.814), temp_mean (1.010). Overall: RED. This explains cold-group conformal coverage failure (exchangeability violated at PSI this large).",
    "drift Figure 9 caption",
)
fix(
    "Why is RED drift expected here?",
    "Why does RED drift cause cold-group coverage failure? PSI analysis shows 5/10 features have RED-level drift (range 1.01-2.57) between training and test batteries, violating the exchangeability assumption required for conformal guarantees. Cold-group batteries (B0041, B0044, B0052), tested at sub-10C, differ most from the room-temperature training set. Empirical result: cold coverage = 74.5% (below 80% target); room and hot over-cover at 100%. This drift-coverage linkage is the mechanistic explanation for the central uncertainty finding.",
    "drift->coverage linkage",
)

# SECTION 10 RAG Llama
fix(
    "Stage 5 goes beyond prediction to generate human-readable explanations",
    "Stage 5 generates human-readable explanations using a local RAG pipeline. Active LLM: Llama 3.1 8B Instruct (4-bit NF4 quantised via bitsandbytes, GPU required). Retrieval: SentenceTransformer (all-MiniLM-L6-v2) embeddings over ChromaDB vector store. Explanation quality assessed via automated rubric (hedging, mechanistic reasoning, source grounding, overconfidence detection): 6/6 pre-computed explanations passed. No retrieval metrics (Recall@k etc.) computed; human evaluation is future work.",
    "RAG Llama section",
)

# SECTION 11 supervisor verdict
fix(
    "Final verdict: CONDITIONAL PASS",
    "Final verdict: CONDITIONAL PASS. System is valid for research use. Key findings: (1) XGBoost multi-seed RMSE = 44.74 +/- 8.38 cycles (best seed: 21.74); (2) Cold conformal coverage = 74.5% -- 5.5 pp below 80% target, caused by RED PSI drift on 5 features; (3) Statistical baseline auto-excluded (unstable across seeds); (4) Survival exploratory only (5 events, no C-index); (5) Anomaly detection unsupervised (no labels). Recommendation: gather more cold-battery data.",
    "supervisor verdict update",
)

# SECTION 13 summary
fix(
    "cold battery conformal coverage being 1.3",
    "The key gap is cold-group conformal coverage at 74.5% vs the 80% target (5.5 pp below), mechanistically explained by RED-level PSI drift on 5/10 features. Room and hot over-cover at 100% with over-wide intervals. Multi-seed XGBoost RMSE: 44.74 +/- 8.38 cycles. Survival and anomaly components are exploratory.",
    "summary 1.3%->5.5%",
)
fix(
    "The only gap is the CV/test RMSE ratio (2.2",
    "In summary: a reproducible cross-battery RUL prediction pipeline with calibrated uncertainty quantification. CV mean RMSE (17.64 cycles) is optimistic relative to multi-seed test RMSE (44.74 +/- 8.38 cycles), reflecting battery-split dependence. Cold-group coverage failure (74.5% vs 80% target) is the central finding, explained by distribution shift (RED PSI on 5 features).",
    "summary CV/test gap",
)

doc.save("Battery_Project_Report.docx")

print(f"Saved. {len(changes)} paragraphs changed.")
for c in changes:
    print(" ", c)
if not_found:
    print(f"\n{len(not_found)} NOT FOUND:")
    for nf in not_found:
        print(" ", nf)
